# core/reports.py
from __future__ import annotations
from typing import List, Dict, Optional, Tuple
import os
from src.database.connection import get_connection
from src.services.inventory import get_product
from decimal import Decimal, ROUND_HALF_UP

D = Decimal

def _fmt2(x: float | int | Decimal | None) -> float:
    if x is None: return 0.0
    return float(D(str(x)).quantize(D("0.01"), rounding=ROUND_HALF_UP))

def kardex_rows(codigo: str, date_from: Optional[str] = None, date_to: Optional[str] = None) -> Tuple[List[Dict], Dict]:
    """
    Devuelve (rows, summary) para un producto.
    rows: dicts ordenados por fecha con columnas:
      fecha, doc_no, tipo, qty_in, qty_out, unit_cost, unit_price, motivo, saldo_qty, saldo_cost, saldo_total
    summary: {saldo_qty, saldo_cost, saldo_total, first_date, last_date, nombre}
    """
    prod = get_product(codigo)
    if not prod:
        raise ValueError(f"Producto '{codigo}' no existe.")
    nombre = prod[1]

    where = ["l.codigo = ?"]
    params = [codigo]
    if date_from: where.append("d.fecha >= ?"); params.append(date_from)
    if date_to:   where.append("d.fecha <= ?"); params.append(date_to)
    where_sql = "WHERE " + " AND ".join(where)

    sql = f"""
    SELECT d.fecha, d.numero, d.tipo, l.qty, l.unit_cost, l.unit_price, l.reason
    FROM document_lines l
    JOIN documents d ON d.id = l.doc_id
    {where_sql}
    ORDER BY d.fecha ASC, d.id ASC, l.id ASC
    """
    rows = []
    with get_connection() as conn:
        cur = conn.cursor()
        for r in cur.execute(sql, params).fetchall():
            fecha, numero, tipo, qty, unit_cost, unit_price, reason = r
            qty = float(qty or 0.0)
            unit_cost = _fmt2(unit_cost) if unit_cost is not None else None
            unit_price = _fmt2(unit_price) if unit_price is not None else None
            rows.append(dict(
                fecha=str(fecha), doc_no=numero or "", tipo=tipo,
                qty=qty, unit_cost=unit_cost, unit_price=unit_price, motivo=reason or ""
            ))

    # Recorrido para saldos y costo promedio móvil
    saldo_qty = 0.0
    saldo_cost = 0.0  # costo promedio vigente
    out_rows: List[Dict] = []
    for r in rows:
        qty = r["qty"]
        if r["tipo"] in ("PURCHASE", "ADJUST+"):
            # entradas usan unit_cost; si viene None, usa saldo_cost actual
            cost = r["unit_cost"] if r["unit_cost"] is not None else (saldo_cost or 0.0)
            # nuevo promedio ponderado
            total_val = saldo_qty * saldo_cost + qty * cost
            saldo_qty = saldo_qty + qty
            saldo_cost = (total_val / saldo_qty) if saldo_qty > 0 else 0.0
            qty_in, qty_out = qty, 0.0
        elif r["tipo"] in ("SALE", "ADJUST-"):
            # salidas valorizadas al costo promedio vigente
            if qty <= 0: raise ValueError("Cantidad de salida inválida.")
            if saldo_qty - qty < -1e-6:
                raise ValueError(f"Stock insuficiente para '{codigo}' en {r['fecha']} (sale {qty}, saldo {saldo_qty}).")
            qty_in, qty_out = 0.0, qty
            saldo_qty = saldo_qty - qty
            # saldo_cost se mantiene
        else:
            # por si a futuro
            qty_in, qty_out = (qty if qty >=0 else 0.0), (abs(qty) if qty < 0 else 0.0)

        out = dict(
            fecha=r["fecha"], doc_no=r["doc_no"], tipo=r["tipo"],
            qty_in=_fmt2(qty_in), qty_out=_fmt2(qty_out),
            unit_cost=_fmt2(r["unit_cost"]) if r["unit_cost"] is not None else _fmt2(saldo_cost),
            unit_price=_fmt2(r["unit_price"]) if r["unit_price"] is not None else None,
            motivo=r["motivo"],
            saldo_qty=_fmt2(saldo_qty), saldo_cost=_fmt2(saldo_cost),
            saldo_total=_fmt2(saldo_qty * saldo_cost)
        )
        out_rows.append(out)

    summary = dict(
        saldo_qty=_fmt2(saldo_qty),
        saldo_cost=_fmt2(saldo_cost),
        saldo_total=_fmt2(saldo_qty * saldo_cost),
        first_date=out_rows[0]["fecha"] if out_rows else None,
        last_date=out_rows[-1]["fecha"] if out_rows else None,
        nombre=nombre
    )
    return out_rows, summary

# ---------- Exportaciones ----------
def export_kardex_xlsx(path: str, codigo: str, date_from: Optional[str]=None, date_to: Optional[str]=None) -> int:
    import openpyxl
    from openpyxl.utils import get_column_letter

    rows, summary = kardex_rows(codigo, date_from, date_to)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kardex"

    ws.append(["Kardex de producto", codigo, summary["nombre"] or "", "", "", "", ""])
    ws.append(["Rango", date_from or "-", date_to or "-", "", "", "", ""])
    ws.append([])

    headers = ["Fecha", "Doc Nº", "Tipo", "Ent.", "Sal.", "Costo u.", "Precio u.", "Motivo", "Saldo", "Costo prom.", "Valorización"]
    ws.append(headers)

    for r in rows:
        ws.append([
            r["fecha"], r["doc_no"], r["tipo"], r["qty_in"], r["qty_out"],
            r["unit_cost"], r["unit_price"] if r["unit_price"] is not None else "",
            r["motivo"], r["saldo_qty"], r["saldo_cost"], r["saldo_total"]
        ])

    ws.append([])
    ws.append(["Totales actuales", "", "", "", "", "", "", "",
               summary["saldo_qty"], summary["saldo_cost"], summary["saldo_total"]])

    # Ajuste de anchos
    for col in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(col)].width = 14

    wb.save(path)
    return len(rows)

def export_kardex_pdf(path: str, codigo: str, date_from: Optional[str]=None, date_to: Optional[str]=None,
                      title: str="Kardex de producto", logo_path: Optional[str]=None) -> int:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader

    rows, summary = kardex_rows(codigo, date_from, date_to)
    c = canvas.Canvas(path, pagesize=A4)
    W, H = A4
    y = H - 2*cm

    if logo_path and os.path.exists(logo_path):
        try:
            img = ImageReader(logo_path)
            c.drawImage(img, 2*cm, y-2*cm, width=2.5*cm, height=2.5*cm, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass

    c.setFont("Helvetica-Bold", 14)
    c.drawString(5*cm, y, title)
    c.setFont("Helvetica", 9)
    c.drawString(5*cm, y-0.6*cm, f"Código: {codigo} — {summary['nombre']}")
    c.drawString(5*cm, y-1.1*cm, f"Rango: {date_from or '-'} a {date_to or '-'}")
    y -= 2*cm

    # Cabeceras
    c.setFont("Helvetica-Bold", 8)
    headers = ["Fecha","Doc","Tipo","Ent.","Sal.","Cst u.","Prc u.","Motivo","Saldo","Cst prom.","Val."]
    x_positions = [1.0, 3.8, 5.2, 6.6, 7.6, 8.6, 9.7, 10.8, 14.0, 15.5, 17.0]  # en cm
    for i, h in enumerate(headers):
        c.drawString(x_positions[i]*cm, y, h)
    y -= 0.35*cm
    c.setFont("Helvetica", 8)

    def new_page():
        nonlocal y
        c.showPage()
        y = H - 2*cm
        c.setFont("Helvetica-Bold", 8)
        for i, h in enumerate(headers):
            c.drawString(x_positions[i]*cm, y, h)
        y -= 0.35*cm
        c.setFont("Helvetica", 8)

    for r in rows:
        if y < 2*cm: new_page()
        vals = [
            r["fecha"][:19], r["doc_no"], r["tipo"], f"{r['qty_in']:.2f}", f"{r['qty_out']:.2f}",
            f"{r['unit_cost']:.2f}", f"{(r['unit_price'] or 0):.2f}" if r["unit_price"] is not None else "",
            (r["motivo"] or "")[:20], f"{r['saldo_qty']:.2f}", f"{r['saldo_cost']:.2f}", f"{r['saldo_total']:.2f}"
        ]
        for i, v in enumerate(vals):
            c.drawString(x_positions[i]*cm, y, str(v))
        y -= 0.35*cm

    # resumen
    if y < 3*cm: new_page()
    y -= 0.2*cm
    c.setFont("Helvetica-Bold", 9)
    c.drawString(1.0*cm, y, f"Saldo: {summary['saldo_qty']:.2f} u | Costo prom.: {summary['saldo_cost']:.2f} | Valorización: {summary['saldo_total']:.2f}")
    c.save()
    return len(rows)
