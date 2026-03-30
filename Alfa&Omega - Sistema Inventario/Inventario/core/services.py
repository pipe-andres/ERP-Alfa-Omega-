# core/services.py
from __future__ import annotations
from typing import List, Tuple, Optional, Dict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import csv
import os

from .database import get_connection
from .audit import log_event
from .documents import get_next_number
from .partners import get_partner_by_code
from . import metrics

D = Decimal
FMT2 = lambda x: float(D(str(x)).quantize(D("0.01"), rounding=ROUND_HALF_UP))

# =========================
# Productos (CRUD / Listado)
# =========================

def get_product(codigo: str) -> Optional[Tuple]:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT codigo, nombre, categoria, precio, cantidad FROM productos WHERE codigo = ?", (codigo,))
        row = cur.fetchone()
        return tuple(row) if row else None

def add_product(codigo: str, nombre: str, categoria: str, precio: float, cantidad: int, user_id: Optional[int] = None):
    if not codigo or not nombre:
        raise ValueError("Código y Nombre son obligatorios.")
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("""
            INSERT INTO productos (codigo, nombre, categoria, precio, cantidad, avg_cost)
            VALUES (?,?,?,?,?,?)
        """, (codigo.strip(), nombre.strip(), (categoria or "").strip(), float(precio), int(cantidad), float(precio)))
        conn.commit()
    if user_id:
        log_event(user_id, "PRODUCT_CREATE", {"codigo": codigo, "nombre": nombre})
    # metrics
    try:
        metrics.increment("products_added", 1)
    except Exception:
        pass

def update_product(codigo: str, nombre: Optional[str] = None, categoria: Optional[str] = None,
                   precio: Optional[float] = None, cantidad: Optional[int] = None, user_id: Optional[int] = None):
    sets, params = [], []
    if nombre is not None:    sets.append("nombre=?");    params.append(nombre.strip())
    if categoria is not None: sets.append("categoria=?"); params.append((categoria or "").strip())
    if precio is not None:    sets.append("precio=?");    params.append(float(precio))
    if cantidad is not None:  sets.append("cantidad=?");  params.append(int(cantidad))
    if not sets:
        return
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("UPDATE productos SET " + ", ".join(sets) + " WHERE codigo=?", params + [codigo])
        if cur.rowcount == 0:
            raise ValueError("Producto no encontrado.")
        conn.commit()
    if user_id:
        log_event(user_id, "PRODUCT_UPDATE", {"codigo": codigo})
    try:
        metrics.increment("products_updated", 1)
    except Exception:
        pass

def delete_product(codigo: str, user_id: Optional[int] = None):
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("DELETE FROM productos WHERE codigo=?", (codigo,))
        if cur.rowcount == 0:
            raise ValueError("Producto no encontrado.")
        conn.commit()
    if user_id:
        log_event(user_id, "PRODUCT_DELETE", {"codigo": codigo})
    try:
        metrics.increment("products_deleted", 1)
    except Exception:
        pass

def count_products(filter_by: Optional[str] = None, q: str = "") -> int:
    where, params = [], []
    if q:
        ql = f"%{q.lower()}%"
        if filter_by == "Código":
            where.append("LOWER(codigo) LIKE ?"); params.append(ql)
        elif filter_by == "Nombre":
            where.append("LOWER(nombre) LIKE ?"); params.append(ql)
        elif filter_by == "Categoría":
            where.append("LOWER(categoria) LIKE ?"); params.append(ql)
        else:
            where.append("(LOWER(codigo) LIKE ? OR LOWER(nombre) LIKE ? OR LOWER(categoria) LIKE ?)")
            params += [ql, ql, ql]
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(f"SELECT COUNT(*) FROM productos {where_sql}", params)
        return int(cur.fetchone()[0])

def list_products_page(page: int = 1, page_size: int = 25,
                       filter_by: Optional[str] = None, q: str = "",
                       order_by: str = "nombre", asc: bool = True) -> List[Tuple]:
    order_allowed = {"codigo","nombre","categoria","precio","cantidad"}
    if order_by not in order_allowed: order_by = "nombre"
    order_sql = "ASC" if asc else "DESC"

    where, params = [], []
    if q:
        ql = f"%{q.lower()}%"
        if filter_by == "Código":
            where.append("LOWER(codigo) LIKE ?"); params.append(ql)
        elif filter_by == "Nombre":
            where.append("LOWER(nombre) LIKE ?"); params.append(ql)
        elif filter_by == "Categoría":
            where.append("LOWER(categoria) LIKE ?"); params.append(ql)
        else:
            where.append("(LOWER(codigo) LIKE ? OR LOWER(nombre) LIKE ? OR LOWER(categoria) LIKE ?)")
            params += [ql, ql, ql]
    where_sql = "WHERE " + " AND ".join(where) if where else ""

    offset = max(0, (int(page)-1) * int(page_size))
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute(f"""
            SELECT codigo, nombre, categoria, precio, cantidad
            FROM productos
            {where_sql}
            ORDER BY {order_by} {order_sql}
            LIMIT ? OFFSET ?
        """, params + [int(page_size), int(offset)])
        return [tuple(row) for row in cur.fetchall()]

def stock_global_sum() -> float:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT COALESCE(SUM(cantidad),0) FROM productos")
        return float(cur.fetchone()[0] or 0)

def low_stock_count(threshold: int = 5) -> int:
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT COUNT(*) FROM productos WHERE cantidad <= ?", (int(threshold),))
        return int(cur.fetchone()[0] or 0)

# =========================
# CSV / Excel
# =========================

def export_products_csv(path: str) -> int:
    with get_connection() as conn, open(path, "w", newline="", encoding="utf-8") as f:
        cur = conn.cursor()
        cur.execute("SELECT codigo, nombre, categoria, precio, cantidad FROM productos ORDER BY nombre ASC")
        rows = [tuple(r) for r in cur.fetchall()]
        w = csv.writer(f)
        w.writerow(["codigo","nombre","categoria","precio","cantidad"])
        for r in rows: w.writerow(r)
        return len(rows)

def import_products_csv(path: str, upsert: bool = True, user_id: Optional[int] = None) -> int:
    count = 0
    with get_connection() as conn, open(path, newline="", encoding="utf-8") as f:
        cur = conn.cursor()
        for row in csv.DictReader(f):
            codigo = (row.get("codigo") or "").strip()
            if not codigo: continue
            nombre = (row.get("nombre") or "").strip()
            categoria = (row.get("categoria") or "").strip()
            precio = float(row.get("precio") or 0.0)
            cantidad = int(float(row.get("cantidad") or 0))
            if upsert:
                cur.execute("""
                    INSERT INTO productos (codigo,nombre,categoria,precio,cantidad,avg_cost)
                    VALUES (?,?,?,?,?,?)
                    ON CONFLICT(codigo) DO UPDATE SET
                        nombre=excluded.nombre,
                        categoria=excluded.categoria,
                        precio=excluded.precio,
                        cantidad=excluded.cantidad,
                        avg_cost=excluded.precio
                """, (codigo, nombre, categoria, precio, cantidad, precio))
            else:
                cur.execute("INSERT INTO productos (codigo,nombre,categoria,precio,cantidad,avg_cost) VALUES (?,?,?,?,?,?)",
                            (codigo, nombre, categoria, precio, cantidad, precio))
            count += 1
        conn.commit()
    if user_id:
        log_event(user_id, "PRODUCT_IMPORT_CSV", {"path": os.path.basename(path), "rows": count})
    return count

def export_products_xlsx(path: str) -> int:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Productos"
    ws.append(["codigo","nombre","categoria","precio","cantidad"])
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT codigo,nombre,categoria,precio,cantidad FROM productos ORDER BY nombre ASC")
        rows = [tuple(r) for r in cur.fetchall()]
        for r in rows: ws.append(list(r))
    wb.save(path)
    return len(rows)

def import_products_xlsx(path: str, upsert: bool = True, user_id: Optional[int] = None) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(path); ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    req = {"codigo","nombre","precio","cantidad"}
    if not req.issubset(idx.keys()):
        raise ValueError(f"Faltan columnas obligatorias: {', '.join(req-idx.keys())}")
    count = 0
    with get_connection() as conn:
        cur = conn.cursor()
        for row in ws.iter_rows(min_row=2, values_only=True):
            codigo = (row[idx["codigo"]] or "").strip()
            if not codigo: continue
            nombre = (row[idx["nombre"]] or "").strip()
            categoria = (row[idx.get("categoria", -1)] or "").strip() if "categoria" in idx else ""
            precio = float(row[idx["precio"]] or 0.0)
            cantidad = int(float(row[idx["cantidad"]] or 0))
            if upsert:
                cur.execute("""
                    INSERT INTO productos (codigo,nombre,categoria,precio,cantidad,avg_cost)
                    VALUES (?,?,?,?,?,?)
                    ON CONFLICT(codigo) DO UPDATE SET
                        nombre=excluded.nombre,
                        categoria=excluded.categoria,
                        precio=excluded.precio,
                        cantidad=excluded.cantidad,
                        avg_cost=excluded.precio
                """, (codigo, nombre, categoria, precio, cantidad, precio))
            else:
                cur.execute("INSERT INTO productos (codigo,nombre,categoria,precio,cantidad,avg_cost) VALUES (?,?,?,?,?,?)",
                            (codigo, nombre, categoria, precio, cantidad, precio))
            count += 1
        conn.commit()
    if user_id:
        log_event(user_id, "PRODUCT_IMPORT_XLSX", {"path": os.path.basename(path), "rows": count})
    return count

# =========================
# PDF Inventario / Stock Bajo
# =========================

def export_inventory_pdf(path: str, logo_path: Optional[str] = None, title: str = "Inventario"):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader

    c = canvas.Canvas(path, pagesize=A4)
    W, H = A4
    y = H - 2*cm

    if logo_path and os.path.exists(logo_path):
        try:
            img = ImageReader(logo_path)
            c.drawImage(img, 2*cm, y-2*cm, width=2.5*cm, height=2.5*cm, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass

    c.setFont("Helvetica-Bold", 14); c.drawString(5*cm, y, title)
    c.setFont("Helvetica", 9); c.drawString(5*cm, y-0.6*cm, f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    y -= 2*cm

    headers = ["Código","Nombre","Categoría","Precio","Cant."]
    x = [1.5, 4.2, 10.0, 15.0, 17.0]
    c.setFont("Helvetica-Bold", 9)
    for i, h in enumerate(headers): c.drawString(x[i]*cm, y, h)
    y -= 0.35*cm
    c.setFont("Helvetica", 9)

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT codigo,nombre,categoria,precio,cantidad FROM productos ORDER BY nombre ASC")
        rows = [tuple(r) for r in cur.fetchall()]
        for cod, nom, cat, pre, can in rows:
            if y < 2*cm:
                c.showPage(); y = H - 2*cm
                c.setFont("Helvetica-Bold", 9)
                for i, h in enumerate(headers): c.drawString(x[i]*cm, y, h)
                y -= 0.35*cm; c.setFont("Helvetica", 9)
            c.drawString(x[0]*cm, y, str(cod))
            c.drawString(x[1]*cm, y, str(nom)[:40])
            c.drawString(x[2]*cm, y, str(cat)[:20])
            c.drawRightString((x[3]+1.8)*cm, y, f"{pre:.2f}")
            c.drawRightString((x[4]+1.2)*cm, y, str(can))
            y -= 0.3*cm
    c.save()

def export_low_stock_pdf(path: str, threshold: int = 5, logo_path: Optional[str] = None, title: str = "Stock Bajo"):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader

    c = canvas.Canvas(path, pagesize=A4)
    W, H = A4
    y = H - 2*cm

    if logo_path and os.path.exists(logo_path):
        try:
            img = ImageReader(logo_path)
            c.drawImage(img, 2*cm, y-2*cm, width=2.5*cm, height=2.5*cm, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass

    c.setFont("Helvetica-Bold", 14); c.drawString(5*cm, y, title)
    c.setFont("Helvetica", 9); c.drawString(5*cm, y-0.6*cm, f"Umbral: ≤ {threshold} — Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    y -= 2*cm

    headers = ["Código","Nombre","Categoría","Precio","Cant."]
    x = [1.5, 4.2, 10.0, 15.0, 17.0]
    c.setFont("Helvetica-Bold", 9)
    for i, h in enumerate(headers): c.drawString(x[i]*cm, y, h)
    y -= 0.35*cm
    c.setFont("Helvetica", 9)

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT codigo,nombre,categoria,precio,cantidad FROM productos WHERE cantidad <= ? ORDER BY cantidad ASC, nombre ASC", (int(threshold),))
        rows = [tuple(r) for r in cur.fetchall()]
        for cod, nom, cat, pre, can in rows:
            if y < 2*cm:
                c.showPage(); y = H - 2*cm
                c.setFont("Helvetica-Bold", 9)
                for i, h in enumerate(headers): c.drawString(x[i]*cm, y, h)
                y -= 0.35*cm; c.setFont("Helvetica", 9)
            c.drawString(x[0]*cm, y, str(cod))
            c.drawString(x[1]*cm, y, str(nom)[:40])
            c.drawString(x[2]*cm, y, str(cat)[:20])
            c.drawRightString((x[3]+1.8)*cm, y, f"{pre:.2f}")
            c.drawRightString((x[4]+1.2)*cm, y, str(can))
            y -= 0.3*cm
    c.save()

# =========================
# Movimientos: Compras/Ventas/Ajustes
# =========================

def post_purchase(numero: str | None, fecha: str | None, items: List[Dict],
                  notas: str | None = None, partner_code: str | None = None,
                  series: str | None = "C01") -> tuple[int, str | None]:
    numero_final = (numero or "").strip() or get_next_number("PURCHASE", (series or "C01").strip().upper())
    partner_id = None
    if partner_code:
        pr = get_partner_by_code(partner_code)
        if not pr:
            raise ValueError(f"Proveedor '{partner_code}' no existe.")
        partner_id = pr[0]
    fecha_ok = fecha or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("INSERT INTO documents (tipo, fecha, numero, notas, partner_id) VALUES (?,?,?,?,?) RETURNING id",
                    ("PURCHASE", fecha_ok, numero_final, notas, partner_id))
        doc_id = int(cur.fetchone()[0])

        for it in items:
            codigo = it["codigo"]; qty = float(it["qty"])
            unit_cost = float(it.get("unit_cost") or 0.0)
            unit_price = it.get("unit_price")
            cur.execute("""INSERT INTO document_lines (doc_id, codigo, qty, unit_cost, unit_price, reason)
                           VALUES (?,?,?,?,?,?)""", (doc_id, codigo, qty, unit_cost, unit_price, ""))
            cur.execute("""INSERT INTO stock_movements (doc_id, codigo, qty, unit_cost, unit_price, tipo, reason, created_at)
                           VALUES (?,?,?,?,?,?,?,?)""", (doc_id, codigo, qty, unit_cost, unit_price, "PURCHASE", "", fecha_ok))
            cur.execute("SELECT cantidad, precio FROM productos WHERE codigo=?", (codigo,))
            row = cur.fetchone()
            if not row:
                cur.execute("INSERT INTO productos (codigo, nombre, categoria, precio, cantidad, avg_cost) VALUES (?,?,?,?,?,?)",
                            (codigo, codigo, "", unit_cost, int(qty), unit_cost))
            else:
                old_qty, old_cost = float(row[0] or 0), float(row[1] or 0)
                new_qty = old_qty + qty
                new_cost = (old_qty * old_cost + qty * unit_cost) / new_qty if new_qty > 0 else unit_cost
                cur.execute("UPDATE productos SET cantidad=?, precio=?, avg_cost=? WHERE codigo=?",
                            (int(new_qty), new_cost, new_cost, codigo))
        conn.commit()

    log_event(None, "DOC_PURCHASE", {"doc_id": doc_id, "numero": numero_final})
    return doc_id, numero_final

def post_sale(numero: str | None, fecha: str | None, items: List[Dict],
              notas: str | None = None, partner_code: str | None = None,
              series: str | None = "V01", allow_negative: bool = False) -> tuple[int, str | None, dict]:
    """
    Registra una VENTA con líneas y afecta stock.
    items: [{codigo, qty, unit_price}]
    Si numero es None, se asigna por serie con get_next_number('SALE', series).
    Retorna (doc_id, numero_final, totals_dict).
    """
    from .settings import get_settings
    s = get_settings()
    tax_rate = float(s.get("tax_rate") or 0.0)
    tax_included = bool(s.get("tax_included"))

    numero_final = (numero or "").strip() or get_next_number("SALE", (series or "V01").strip().upper())
    partner_id = None
    if partner_code:
        pr = get_partner_by_code(partner_code)
        if not pr:
            raise ValueError(f"Cliente '{partner_code}' no existe.")
        partner_id = pr[0]
    fecha_ok = fecha or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # cálculo subtotales (ventas)
    from decimal import Decimal, ROUND_HALF_UP
    D = Decimal
    q2 = lambda x: D(str(x)).quantize(D("0.01"), rounding=ROUND_HALF_UP)

    subtotal = D("0.00")
    for it in items:
        qty = D(str(it.get("qty") or 0))
        price = D(str(it.get("unit_price") or 0))
        subtotal += (qty * price)

    # impuestos
    neto = subtotal
    tax = D("0.00")
    total = subtotal
    if tax_rate > 0:
        r = D(str(tax_rate))
        if tax_included:
            base = (subtotal / (D("1") + r)).quantize(D("0.01"), rounding=ROUND_HALF_UP)
            neto = base
            tax = (subtotal - base).quantize(D("0.01"), rounding=ROUND_HALF_UP)
            total = subtotal
        else:
            tax = (subtotal * r).quantize(D("0.01"), rounding=ROUND_HALF_UP)
            total = (subtotal + tax).quantize(D("0.01"), rounding=ROUND_HALF_UP)
            neto = subtotal

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("INSERT INTO documents (tipo, fecha, numero, notas, partner_id) VALUES (?,?,?,?,?) RETURNING id",
                    ("SALE", fecha_ok, numero_final, notas, partner_id))
        doc_id = int(cur.fetchone()[0])

        for it in items:
            codigo = it["codigo"]; qty = float(it["qty"])
            unit_price = float(it.get("unit_price") or 0.0)
            # línea
            cur.execute("""INSERT INTO document_lines (doc_id, codigo, qty, unit_cost, unit_price, reason)
                           VALUES (?,?,?,?,?,?)""", (doc_id, codigo, qty, None, unit_price, ""))

            # salida de stock valuada al costo promedio vigente (precio almacenado en productos):
            cur.execute("SELECT cantidad, precio FROM productos WHERE codigo=?", (codigo,))
            row = cur.fetchone()
            if not row:
                raise ValueError(f"Producto '{codigo}' no existe.")
            old_qty, avg_cost = float(row[0] or 0), float(row[1] or 0)
            if (old_qty - qty) < -1e-6 and not allow_negative:
                raise ValueError(f"Stock insuficiente para '{codigo}'. Disponible: {old_qty}, requerido: {qty}")
            new_qty = old_qty - qty
            cur.execute("UPDATE productos SET cantidad=? WHERE codigo=?", (int(new_qty), codigo))
            # registrar movimiento con unit_cost = avg_cost (para kardex)
            cur.execute("""INSERT INTO stock_movements (doc_id, codigo, qty, unit_cost, unit_price, tipo, reason, created_at)
                           VALUES (?,?,?,?,?,?,?,?)""", (doc_id, codigo, -qty, avg_cost, unit_price, "SALE", "", fecha_ok))
        conn.commit()

    log_event(None, "DOC_SALE", {"doc_id": doc_id, "numero": numero_final})

    return doc_id, numero_final, {
        "subtotal": float(q2(subtotal)),
        "neto": float(q2(neto)),
        "impuesto": float(q2(tax)),
        "total": float(q2(total)),
        "tax_included": tax_included,
        "tax_rate": float(tax_rate),
    }

def post_adjustment(fecha: Optional[str], items: List[Dict], notas: Optional[str] = None) -> int:
    fecha_ok = fecha or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("INSERT INTO documents (tipo, fecha, numero, notas, partner_id) VALUES (?,?,?,?,NULL) RETURNING id",
                    ("ADJUST", fecha_ok, None, notas))
        doc_id = int(cur.fetchone()[0])

        for it in items:
            codigo = it["codigo"]; qty = float(it["qty"])
            reason = (it.get("reason") or "").strip()
            unit_cost_in = it.get("unit_cost")
            cur.execute("""INSERT INTO document_lines (doc_id, codigo, qty, unit_cost, unit_price, reason)
                           VALUES (?,?,?,?,?,?)""", (doc_id, codigo, qty, unit_cost_in, None, reason))
            cur.execute("SELECT cantidad, precio FROM productos WHERE codigo=?", (codigo,))
            row = cur.fetchone()
            if not row:
                raise ValueError(f"Producto '{codigo}' no existe.")
            old_qty, avg_cost = float(row[0] or 0), float(row[1] or 0)
            if qty >= 0:
                if unit_cost_in is None: unit_cost_in = avg_cost
                new_qty = old_qty + qty
                new_cost = (old_qty * avg_cost + qty * float(unit_cost_in)) / new_qty if new_qty > 0 else float(unit_cost_in)
                cur.execute("UPDATE productos SET cantidad=?, precio=?, avg_cost=? WHERE codigo=?", (int(new_qty), new_cost, new_cost, codigo))
                cur.execute("""INSERT INTO stock_movements (doc_id, codigo, qty, unit_cost, unit_price, tipo, reason, created_at)
                               VALUES (?,?,?,?,?,?,?,?)""", (doc_id, codigo, qty, float(unit_cost_in), None, "ADJUST+", reason, fecha_ok))
            else:
                if (old_qty + qty) < -1e-6:
                    raise ValueError(f"Stock insuficiente para ajustar '{codigo}'. Disponible: {old_qty}, salida: {-qty}")
                new_qty = old_qty + qty
                cur.execute("UPDATE productos SET cantidad=? WHERE codigo=?", (int(new_qty), codigo))
                cur.execute("""INSERT INTO stock_movements (doc_id, codigo, qty, unit_cost, unit_price, tipo, reason, created_at)
                               VALUES (?,?,?,?,?,?,?,?)""", (doc_id, codigo, qty, avg_cost, None, "ADJUST-", reason, fecha_ok))
        conn.commit()
    log_event(None, "DOC_ADJUST", {"doc_id": doc_id, "items": len(items)})
    return doc_id
